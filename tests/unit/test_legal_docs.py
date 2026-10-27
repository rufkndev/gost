import io

import docx

from app.config import Settings
from app.services.legal_docs import DOCS, all_files


def test_files_filled_with_requisites() -> None:
    s = Settings(
        _env_file=None,
        ip_name="Иванов Иван Иванович",
        inn="123456789012",
        ogrnip="312345678901234",
        contact_email="ip@example.ru",
        ip_registration="зарегистрирован ИФНС № 1 01.01.2020",
    )
    files = dict(all_files(s))
    assert set(files) == {d.filename for d in DOCS}
    for data in files.values():
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
        assert "{{" not in text  # все реквизиты подставлены
        assert "ЧЕРНОВИК" in text
    offer = "\n".join(
        p.text for p in docx.Document(io.BytesIO(files["Публичная оферта.docx"])).paragraphs
    )
    assert "Иванов Иван Иванович" in offer and "123456789012" in offer


def test_missing_requisites_stay_visible() -> None:
    files = dict(all_files(Settings(_env_file=None)))
    text = "\n".join(
        p.text for p in docx.Document(io.BytesIO(files["Публичная оферта.docx"])).paragraphs
    )
    assert "{{IP_NAME}}" in text  # незаполненное видно сразу, а не пустое место

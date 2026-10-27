"""Юридические документы как файлы .docx (для отправки в боте и для владельца).

Источник — docs/legal/*.md с плейсхолдерами {{…}}; реквизиты подставляются из конфига.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from functools import lru_cache

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from app.config import ROOT_DIR, Settings

LEGAL_DIR = ROOT_DIR / "docs" / "legal"


def placeholders(s: Settings) -> dict[str, str]:
    """Значения {{…}}; незаполненный реквизит остаётся плейсхолдером — его сразу видно."""
    return {
        "IP_NAME": s.ip_name or "{{IP_NAME}}",
        "INN": s.inn or "{{INN}}",
        "OGRNIP": s.ogrnip or "{{OGRNIP}}",
        "IP_REGISTRATION": s.ip_registration or "{{IP_REGISTRATION}}",
        "EMAIL": s.contact_email or "{{EMAIL}}",
        "SUPPORT": s.support_contact or s.contact_email or "{{SUPPORT}}",
        "PRICE": str(s.price_rub),
        "TERMS_VERSION": s.terms_version,
        "RETENTION_UNPAID_HOURS": str(s.retention_unpaid_hours),
        "RETENTION_PAID_HOURS": str(s.retention_paid_hours),
        "RETENTION_COMPLAINT_DAYS": str(s.retention_complaint_days),
    }


def fill(text: str, values: dict[str, str]) -> str:
    return re.sub(r"\{\{(\w+)\}\}", lambda m: values.get(m.group(1), m.group(0)), text)


@dataclass(frozen=True)
class LegalDoc:
    key: str
    source: str
    filename: str


DOCS = (
    LegalDoc("offer", "offer.md", "Публичная оферта.docx"),
    LegalDoc("privacy", "privacy.md", "Политика обработки персональных данных.docx"),
    LegalDoc("terms", "terms.md", "Правила использования.docx"),
)


def _runs(paragraph: object, text: str) -> None:
    """Текст с **жирными** фрагментами."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            run = paragraph.add_run(part)  # type: ignore[attr-defined]
            run.bold = i % 2 == 1


def markdown_to_docx(md: str) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
    for line in md.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(s[2:].strip())
            r.bold = True
            r.font.size = Pt(14)
        elif s.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            r = p.add_run(s[3:].strip())
            r.bold = True
        elif s.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            _runs(p, "– " + s[2:])
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _runs(p, s)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@lru_cache(maxsize=8)
def _cached(key: str, values: tuple[tuple[str, str], ...]) -> bytes:
    doc = next(d for d in DOCS if d.key == key)
    md = (LEGAL_DIR / doc.source).read_text(encoding="utf-8")
    return markdown_to_docx(fill(md, dict(values)))


def all_files(settings: Settings) -> list[tuple[str, bytes]]:
    values = tuple(sorted(placeholders(settings).items()))
    return [(d.filename, _cached(d.key, values)) for d in DOCS]
